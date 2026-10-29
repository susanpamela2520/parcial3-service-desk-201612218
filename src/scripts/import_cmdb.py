"""Importa los CI de ejemplo de data/CMDB.xlsx (hoja CMDB, encabezados B4:U4, filas 5–6).

Repetible: la llave es el código CMDB-<nombre normalizado>. Los textos se recortan (el Excel trae espacios
iniciales) y se conserva el responsable original como texto. Reglas documentadas en docs/MODELO.md §4:
- I «Relaciones» = «Base de Datos1» -> relación DEPENDE_DE hacia un CI 'Base de Datos1' (clase SW-BD),
  que se crea porque es referenciado pero no tiene fila propia (observación).
- K6 «Ubicación Física» = «Servidor1» no es un lugar: se interpreta como relación ALOJADO_EN Servidor1 y
  la ubicación física del CI se toma de Servidor1 (observación).
- Columnas S y T traen texto mal codificado en el encabezado («Configuraci√≥n»): se mapean por posición.
"""
import os
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from app import config
from app.db import SessionLocal

RUTA = Path(os.getenv("CMDB_PATH", str(config.BASE_DIR / "data" / "CMDB.xlsx")))
CLASE_POR_FILA = {5: "HW-SRV", 6: "SW-APP"}  # decisión documentada (B5 Servidor1, B6 Aplicación)


def limpio(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def codigo(nombre: str) -> str:
    s = "".join(c for c in unicodedata.normalize("NFD", nombre) if unicodedata.category(c) != "Mn")
    return "CMDB-" + re.sub(r"[^A-Z0-9]+", "-", s.upper()).strip("-")


def fecha(v):
    v = limpio(v)
    return date.fromisoformat(v) if v else None


def main() -> int:
    if not RUTA.exists():
        print(f"[cmdb.xlsx] no existe {RUTA}")
        return 1
    ws = load_workbook(RUTA, data_only=True)["CMDB"]
    db = SessionLocal()
    obs = []
    resp = db.execute(text("SELECT id FROM usuario WHERE rol='administrador' AND activo ORDER BY id LIMIT 1")).scalar()
    if resp is None:
        print("[cmdb.xlsx] ejecute seed_demo antes (se necesita un usuario responsable de referencia)")
        return 1
    clase = dict(db.execute(text("SELECT codigo, id FROM ci_clase")).all())
    ubic = dict(db.execute(text("SELECT nombre, id FROM ubicacion")).all())
    ids, creados = {}, 0

    def upsert(fila: dict) -> int:
        nonlocal creados
        cid = db.execute(text("SELECT id FROM ci WHERE codigo=:c"), {"c": fila["codigo"]}).scalar()
        if cid:
            return cid
        creados += 1
        return db.execute(text("""INSERT INTO ci (codigo, nombre, clase_id, descripcion, numero_serie, version,
            fecha_adquisicion, estado_actual, ubicacion_id, responsable_id, responsable_original, fecha_cambio,
            descripcion_cambio, documentacion, enlaces_incidentes, nivel_seguridad, cumplimiento, estado_configuracion,
            numero_licencia, fecha_vencimiento, lote) VALUES (:codigo, :nombre, :clase_id, :descripcion, :numero_serie,
            :version, :fecha_adquisicion, :estado_actual, :ubicacion_id, :responsable_id, :responsable_original,
            :fecha_cambio, :descripcion_cambio, :documentacion, :enlaces_incidentes, :nivel_seguridad, :cumplimiento,
            :estado_configuracion, :numero_licencia, :fecha_vencimiento, 'CMDB.xlsx') RETURNING id"""), fila).scalar()

    filas = {}
    for r in (5, 6):
        v = {c: limpio(ws[f"{c}{r}"].value) for c in "BCDEFGHIJKLMNOPQRSTU"}
        filas[r] = v
    srv_ubic = ubic.get(filas[5]["K"])
    for r, v in filas.items():
        ub = ubic.get(v["K"])
        if ub is None:
            obs.append(f"K{r}='{v['K']}' no es una ubicación física; se usa la ubicación de Servidor1 y se crea ALOJADO_EN")
            ub = srv_ubic
        ids[v["B"]] = upsert({
            "codigo": codigo(v["B"]), "nombre": v["B"], "clase_id": clase[CLASE_POR_FILA[r]], "descripcion": v["D"],
            "numero_serie": v["E"], "version": v["F"], "fecha_adquisicion": fecha(v["G"]), "estado_actual": v["H"],
            "ubicacion_id": ub, "responsable_id": resp, "responsable_original": v["L"], "fecha_cambio": fecha(v["M"]),
            "descripcion_cambio": v["N"], "documentacion": v["O"], "enlaces_incidentes": v["P"],
            "nivel_seguridad": v["Q"], "cumplimiento": v["R"], "estado_configuracion": v["S"],
            "numero_licencia": v["T"], "fecha_vencimiento": fecha(v["U"])})
    for r, v in filas.items():
        if v["I"]:
            if v["I"] not in ids:
                obs.append(f"I{r}='{v['I']}' es referenciado pero no tiene fila propia: se crea como CI de clase SW-BD")
                ids[v["I"]] = upsert({k: None for k in ("descripcion", "numero_serie", "version", "fecha_adquisicion",
                                      "responsable_original", "fecha_cambio", "descripcion_cambio", "documentacion",
                                      "enlaces_incidentes", "nivel_seguridad", "cumplimiento", "estado_configuracion",
                                      "numero_licencia", "fecha_vencimiento")} | {
                    "codigo": codigo(v["I"]), "nombre": v["I"], "clase_id": clase["SW-BD"], "estado_actual": "Activo",
                    "ubicacion_id": srv_ubic, "responsable_id": resp})
            db.execute(text("INSERT INTO ci_relacion (origen_id, destino_id, tipo, lote) VALUES (:o,:d,'DEPENDE_DE','CMDB.xlsx') "
                            "ON CONFLICT DO NOTHING"), {"o": ids[v["B"]], "d": ids[v["I"]]})
        if v["K"] in ids and v["K"] != v["B"]:
            db.execute(text("INSERT INTO ci_relacion (origen_id, destino_id, tipo, lote) VALUES (:o,:d,'ALOJADO_EN','CMDB.xlsx') "
                            "ON CONFLICT DO NOTHING"), {"o": ids[v["B"]], "d": ids[v["K"]]})
        if v["J"]:
            obs.append(f"J{r} «Padres/Hijos» tiene valor '{v['J']}' (no se interpreta automáticamente)")
    db.commit()
    print(f"[cmdb.xlsx] CI creados: {creados}; CI de referencia: {sorted(ids)}")
    for o in obs:
        print(f"  - {o}")
    db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
